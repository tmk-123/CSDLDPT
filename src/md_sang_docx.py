"""Chuyển đổi file báo cáo Markdown (Bao_cao_BTL.md) sang định dạng Word (.docx)
chuẩn học thuật cao cấp, định dạng bảng biểu, hình ảnh, công thức toán và phân cấp đề mục chuyên nghiệp.

Đầu vào : reports/Bao_cao_BTL.md (hoặc Bao_cao_BTL.md)
Đầu ra  : reports/Bao_cao_BTL_HCSDLDPT.docx và Bao_cao_BTL.docx
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

GOC = Path(__file__).resolve().parent.parent
FILE_MD = GOC / "reports" / "Bao_cao_BTL.md"
if not FILE_MD.exists():
    FILE_MD = GOC / "Bao_cao_BTL.md"

OUT_DOCX_REPORT = GOC / "reports" / "Bao_cao_BTL_HCSDLDPT.docx"
OUT_DOCX_ROOT = GOC / "Bao_cao_BTL.docx"

# Bảng màu chuẩn học thuật (Academic Navy & Slate)
MAU_NAVY = RGBColor(27, 54, 93)       # #1B365D - Tiêu đề cấp 1
MAU_DEEP_BLUE = RGBColor(44, 82, 130) # #2C5282 - Tiêu đề cấp 2
MAU_MID_BLUE = RGBColor(43, 108, 176) # #2B6CB0 - Tiêu đề cấp 3
MAU_DARK = RGBColor(45, 55, 72)       # #2D3748 - Tiêu đề cấp 4 & Bullet
MAU_MUTED = RGBColor(113, 128, 150)   # #718096 - Chú thích ảnh
MAU_TEXT = RGBColor(26, 32, 44)       # #1A202C - Nội dung

HEX_NAVY_HEADER = "2C5282"
HEX_ROW_ALT = "F8FAFC"
HEX_BORDER = "CBD5E0"
HEX_CALLOUT_IMPORTANT_BG = "FFFBEB"
HEX_CALLOUT_IMPORTANT_BORDER = "D97706"
HEX_CALLOUT_NOTE_BG = "F0F9FF"
HEX_CALLOUT_NOTE_BORDER = "0284C7"
HEX_CODE_BG = "F1F5F9"


def clean_latex_to_text(text: str) -> str:
    """Chuyển đổi các ký hiệu LaTeX thông dụng sang unicode toán học dễ đọc trong Word."""
    replacements = [
        (r"\\approx", "≈"),
        (r"\\le", "≤"),
        (r"\\ge", "≥"),
        (r"\\pm", "±"),
        (r"\\times", "×"),
        (r"\\cdot", "·"),
        (r"\\in", "∈"),
        (r"\\mathbb\{R\}", "ℝ"),
        (r"\\sum", "∑"),
        (r"\\sqrt", "√"),
        (r"\\mu", "μ"),
        (r"\\sigma", "σ"),
        (r"\\lambda", "λ"),
        (r"\\pi", "π"),
        (r"\\to", "→"),
        (r"\\rightarrow", "→"),
        (r"\\longrightarrow", "→"),
        (r"\\iff", "⇔"),
        (r"\\forall", "∀"),
        (r"\\dots", "..."),
        (r"\\operatorname\{([^\}]+)\}", r"\1"),
        (r"\\text\{([^\}]+)\}", r"\1"),
        (r"\\mathbf\{([^\}]+)\}", r"\1"),
        (r"\\boldsymbol\{([^\}]+)\}", r"\1"),
        (r"\\mathcal\{([^\}]+)\}", r"\1"),
        (r"\\dim", "dim"),
        (r"\\log", "log"),
        (r"\\min", "min"),
        (r"\\max", "max"),
        (r"\\sgn", "sgn"),
        (r"\\cos", "cos"),
        (r"\\ell_([12])", r"ℓ\1"),
        (r"\{,\}", ","),
    ]
    res = text
    for pat, rep in replacements:
        res = re.sub(pat, rep, res)
    return res


def set_cell_background(cell, hex_color: str):
    """Đặt màu nền cho cell trong Word table."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề trong (padding) cho cell (đơn vị dxa, 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def set_table_borders(table, hex_color="CBD5E0", sz="4"):
    """Đặt viền bảng màu xám thanh nhã."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def add_page_number_to_footer(footer_p):
    """Thêm trường số trang tự động vào Footer của Word."""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = MAU_MUTED
    
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def format_runs_in_paragraph(p, text: str, default_font="Times New Roman", default_size=12, default_color=MAU_TEXT):
    """Phân tách các định dạng markdown inline: **bold**, *italic*, `code`, $math$."""
    text = clean_latex_to_text(text)
    pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^\$]+?\$|<[^>]+>)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        run = p.add_run()
        run.font.name = default_font
        run.font.size = Pt(default_size)
        run.font.color.rgb = default_color
        
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 1)
            run.font.color.rgb = RGBColor(199, 37, 78)
        elif part.startswith("$") and part.endswith("$") and len(part) >= 2:
            math_content = part[1:-1].strip()
            run.text = math_content
            run.font.name = "Cambria Math"
            run.italic = True
        elif part.startswith("<") and part.endswith(">"):
            if part.lower() == "<br>":
                run.text = "\n"
        else:
            run.text = part


def tao_trang_bia(doc):
    """Tạo trang bìa trang trọng theo chuẩn Học viện Công nghệ Bưu chính Viễn thông."""
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)
    r1 = p_header.add_run("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = "Times New Roman"
    r1.font.color.rgb = MAU_NAVY
    
    r2 = p_header.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = "Times New Roman"
    r2.font.color.rgb = MAU_NAVY
    
    r3 = p_header.add_run("———————————")
    r3.font.size = Pt(11)
    r3.font.color.rgb = MAU_MUTED
    
    # Khoảng cách giữa
    for _ in range(3):
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_t1 = p_title.add_run("BÁO CÁO BÀI TẬP LỚN\n")
    r_t1.bold = True
    r_t1.font.size = Pt(18)
    r_t1.font.name = "Times New Roman"
    r_t1.font.color.rgb = MAU_NAVY
    
    r_t2 = p_title.add_run("HỆ CƠ SỞ DỮ LIỆU ĐA PHƯƠNG TIỆN")
    r_t2.bold = True
    r_t2.font.size = Pt(16)
    r_t2.font.name = "Times New Roman"
    r_t2.font.color.rgb = MAU_DEEP_BLUE
    
    for _ in range(2):
        doc.add_paragraph()
        
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_s1 = p_sub.add_run("ĐỀ TÀI: XÂY DỰNG HỆ CSDL LƯU TRỮ VÀ TÌM KIẾM\nTIẾNG NHẠC CỤ BỘ DÂY DỰA TRÊN NỘI DUNG (CBAR)")
    r_s1.bold = True
    r_s1.font.size = Pt(14)
    r_s1.font.name = "Times New Roman"
    r_s1.font.color.rgb = MAU_NAVY
    
    # Khung thông tin tác giả
    for _ in range(3):
        doc.add_paragraph()
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(5.8)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_NAVY_HEADER}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p_info = cell.paragraphs[0]
    p_info.paragraph_format.line_spacing = 1.3
    p_info.paragraph_format.space_after = Pt(2)
    
    items = [
        ("Học phần:", " Hệ cơ sở dữ liệu đa phương tiện"),
        ("Sinh viên thực hiện:", " Trần Mạnh Khiêm"),
        ("Nhiệm vụ báo cáo:", ""),
        ("  • Nhiệm vụ 1:", " Xây dựng bộ dữ liệu âm thanh nhạc cụ bộ dây và phân tích đặc điểm giống, khác nhau."),
        ("  • Nhiệm vụ 2:", " Xây dựng bộ đặc trưng nhận diện và thiết kế mô hình tính độ tương đồng cho hệ thống CBAR.")
    ]
    for lbl, val in items:
        p_row = cell.add_paragraph() if lbl != items[0][0] else p_info
        p_row.paragraph_format.line_spacing = 1.25
        p_row.paragraph_format.space_after = Pt(3)
        r_lbl = p_row.add_run(lbl)
        r_lbl.bold = True
        r_lbl.font.size = Pt(11.5)
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.color.rgb = MAU_NAVY
        if val:
            r_val = p_row.add_run(val)
            r_val.font.size = Pt(11.5)
            r_val.font.name = "Times New Roman"
            r_val.font.color.rgb = MAU_TEXT
            
    for _ in range(4):
        doc.add_paragraph()
        
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_foot.add_run("Hà Nội – 2026")
    r_f.bold = True
    r_f.font.size = Pt(12)
    r_f.font.name = "Times New Roman"
    r_f.font.color.rgb = MAU_NAVY
    
    doc.add_page_break()


def render_markdown_to_docx(md_path: Path, out_path: Path):
    print(f">>> Đang đọc file Markdown: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = docx.Document()
    
    # Thiết lập lề trang chuẩn đồ án (Trái 3.0cm, Phải 2.0cm, Trên 2.0cm, Dưới 2.0cm)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.79)      # ~2.0 cm
    sec.bottom_margin = Inches(0.79)   # ~2.0 cm
    sec.left_margin = Inches(1.18)     # ~3.0 cm
    sec.right_margin = Inches(0.79)    # ~2.0 cm
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)
    sec.different_first_page_header_footer = True
    
    # Thêm số trang vào footer từ trang 2 trở đi
    footer = sec.footer
    footer_p = footer.paragraphs[0]
    add_page_number_to_footer(footer_p)
        
    # Tạo trang bìa
    tao_trang_bia(doc)
    
    i = 0
    total_lines = len(lines)
    
    # Bỏ qua phần header bìa trong file MD (từ dòng 1 đến dấu --- đầu tiên)
    while i < total_lines:
        line_clean = lines[i].strip()
        if line_clean == "---":
            i += 1
            break
        i += 1
        
    print(f">>> Đang xử lý nội dung chính từ dòng {i}/{total_lines}...")
    
    in_code_block = False
    code_lines = []
    
    in_table = False
    table_lines = []
    
    in_callout = False
    callout_type = "NOTE"
    callout_lines = []
    
    in_math_block = False
    math_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        rows_data = []
        for tl in table_lines:
            parts = [c.strip() for c in tl.split("|")]
            if len(parts) >= 3:
                cols = parts[1:-1]
                if all(re.match(r'^:?-+:?$', c) for c in cols):
                    continue
                rows_data.append(cols)
                
        if not rows_data:
            table_lines = []
            in_table = False
            return
            
        num_cols = max(len(r) for r in rows_data)
        num_rows = len(rows_data)
        
        tbl = doc.add_table(rows=num_rows, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        set_table_borders(tbl, hex_color=HEX_BORDER)
        
        for r_idx, row in enumerate(rows_data):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                cell_text = row[c_idx] if c_idx < len(row) else ""
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                if is_header:
                    set_cell_background(cell, HEX_NAVY_HEADER)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, HEX_ROW_ALT)
                        
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if is_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    clean_txt = clean_latex_to_text(cell_text)
                    r = p.add_run(clean_txt)
                    r.bold = True
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    format_runs_in_paragraph(p, cell_text, default_size=10, default_color=MAU_TEXT)
                    
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)
        table_lines = []
        in_table = False

    def flush_callout():
        nonlocal in_callout, callout_lines, callout_type
        if not callout_lines:
            in_callout = False
            return
            
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.0)
        cell = tbl.cell(0, 0)
        
        border_col = HEX_CALLOUT_IMPORTANT_BORDER if callout_type == "IMPORTANT" else HEX_CALLOUT_NOTE_BORDER
        bg_col = HEX_CALLOUT_IMPORTANT_BG if callout_type == "IMPORTANT" else HEX_CALLOUT_NOTE_BG
        
        set_cell_background(cell, bg_col)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_col}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        for idx, cl in enumerate(callout_lines):
            p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            format_runs_in_paragraph(p, cl, default_size=11, default_color=MAU_DARK)
            
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)
        callout_lines = []
        in_callout = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
            
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.0)
        cell = tbl.cell(0, 0)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(0)
        
        code_text = "\n".join(code_lines)
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    def flush_math():
        nonlocal in_math_block, math_lines
        if not math_lines:
            in_math_block = False
            return
        formula = " ".join([l.strip() for l in math_lines if l.strip()])
        if formula:
            clean_formula = clean_latex_to_text(formula)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            
            r = p.add_run(clean_formula)
            r.font.name = "Cambria Math"
            r.font.size = Pt(11.5)
            r.italic = True
            r.font.color.rgb = MAU_NAVY
            
        math_lines = []
        in_math_block = False

    images_added = 0
    tables_added = 0

    while i < total_lines:
        line = lines[i].rstrip("\r\n")
        stripped = line.strip()
        
        # Phòng vệ: Nếu gặp Tiêu đề, Bảng, Ảnh, hoặc Ngắt dòng mà đang trong khối math, tự động flush
        if (stripped.startswith("#") or (stripped.startswith("|") and stripped.endswith("|")) or 
            stripped.startswith("![") or stripped == "---") and in_math_block:
            flush_math()

        # 1. Khối code block ```
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table: flush_table()
                if in_callout: flush_callout()
                if in_math_block: flush_math()
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 2. Khối công thức hiển thị $$
        if stripped == "$$":
            if in_math_block:
                flush_math()
            else:
                if in_table: flush_table()
                if in_callout: flush_callout()
                in_math_block = True
                math_lines = []
            i += 1
            continue
            
        if in_math_block:
            math_lines.append(stripped)
            i += 1
            continue

        # 3. Khối Callout > [!NOTE], > [!IMPORTANT]
        if stripped.startswith(">"):
            if in_table: flush_table()
            if in_math_block: flush_math()
            sub = stripped[1:].strip()
            if sub.startswith("[!IMPORTANT]"):
                in_callout = True
                callout_type = "IMPORTANT"
                callout_lines = []
                i += 1
                continue
            elif sub.startswith("[!NOTE]") or sub.startswith("[!TIP]"):
                in_callout = True
                callout_type = "NOTE"
                callout_lines = []
                i += 1
                continue
            elif in_callout:
                callout_lines.append(sub)
                i += 1
                continue
            else:
                in_callout = True
                callout_type = "NOTE"
                callout_lines = [sub]
                i += 1
                continue
        else:
            if in_callout:
                flush_callout()

        # 4. Bảng Markdown (| ... |)
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                if in_math_block: flush_math()
                in_table = True
                table_lines = []
                tables_added += 1
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # 5. Dòng phân cách ---
        if stripped == "---":
            # Kiểm tra xem dòng tiếp theo (bỏ qua dòng trống) có phải là CHƯƠNG không
            next_idx = i + 1
            while next_idx < total_lines and not lines[next_idx].strip():
                next_idx += 1
            if next_idx < total_lines and (lines[next_idx].strip().startswith("## CHƯƠNG") or lines[next_idx].strip().startswith("## Tài liệu")):
                doc.add_page_break()
            else:
                p_hr = doc.add_paragraph()
                p_hr.paragraph_format.space_before = Pt(6)
                p_hr.paragraph_format.space_after = Pt(6)
                r_hr = p_hr.add_run("—" * 45)
                r_hr.font.color.rgb = RGBColor(203, 213, 225)
            i += 1
            continue

        # 6. Hình ảnh: ![caption](path) có thể kèm chú thích cùng dòng
        img_match = re.search(r'!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if img_match:
            alt_text = img_match.group(1).strip()
            img_rel_path = img_match.group(2).strip()
            
            # Kiểm tra xem có caption liền sau cùng dòng không: ví dụ *Hình 2.2: ...*
            after_img = stripped[img_match.end():].strip()
            caption_on_line = None
            if after_img.startswith("*") and after_img.endswith("*"):
                caption_on_line = after_img[1:-1].strip()
            
            # Tìm đường dẫn thực tế của ảnh
            img_path = None
            candidate_paths = [
                GOC / "reports" / img_rel_path,
                GOC / img_rel_path,
                GOC / "reports" / "hinh" / Path(img_rel_path).name,
                GOC / "hinh" / Path(img_rel_path).name,
                GOC / "test" / Path(img_rel_path).name,
            ]
            for cp in candidate_paths:
                if cp.exists():
                    img_path = cp
                    break
                    
            if img_path:
                try:
                    # Chèn ảnh trực tiếp và căn giữa
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(10)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img.add_picture(str(img_path.resolve()), width=Inches(5.8))
                    images_added += 1
                    print(f"  [ĐÃ CHÈN ẢNH {images_added}] {img_path.name}")
                    
                    # Nếu có caption cùng dòng thì chèn luôn
                    if caption_on_line:
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(2)
                        p_cap.paragraph_format.space_after = Pt(10)
                        r_cap = p_cap.add_run(caption_on_line)
                        r_cap.font.name = "Times New Roman"
                        r_cap.font.size = Pt(10)
                        r_cap.font.color.rgb = MAU_MUTED
                        r_cap.italic = True
                except Exception as e:
                    print(f"  [CẢNH BÁO] Không thể chèn ảnh {img_path}: {e}")
                    
            i += 1
            continue

        # 7. Dòng chú thích hình ảnh riêng biệt: *Hình X.Y: ...*
        if (stripped.startswith("*Hình ") or stripped.startswith("*Bảng ")) and stripped.endswith("*"):
            cap_text = stripped[1:-1].strip()
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run(cap_text)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(10)
            r_cap.font.color.rgb = MAU_MUTED
            r_cap.italic = True
            i += 1
            continue

        # 8. Tiêu đề các cấp: #, ##, ###, ####, #####
        if stripped.startswith("#"):
            hashes = len(stripped) - len(stripped.lstrip("#"))
            h_text = stripped.lstrip("#").strip()
            h_clean = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', h_text)
            
            p_h = doc.add_paragraph()
            if hashes == 1 or (hashes == 2 and (h_clean.startswith("CHƯƠNG") or h_clean.startswith("MỤC LỤC"))):
                # Tiêu đề Chương (Heading 1)
                p_h.paragraph_format.space_before = Pt(18)
                p_h.paragraph_format.space_after = Pt(6)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(h_clean)
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(15)
                r.font.color.rgb = MAU_NAVY
            elif hashes == 2 or (hashes == 3 and re.match(r'^\d+\.\d+\.', h_clean)):
                # Tiêu đề mục lớn 1.1, 2.1... (Heading 2)
                p_h.paragraph_format.space_before = Pt(13)
                p_h.paragraph_format.space_after = Pt(4)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(h_clean)
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(13.5)
                r.font.color.rgb = MAU_DEEP_BLUE
            elif hashes == 3 or (hashes == 4 and re.match(r'^\d+\.\d+\.\d+\.', h_clean)):
                # Tiêu đề mục con 1.1.1... (Heading 3)
                p_h.paragraph_format.space_before = Pt(10)
                p_h.paragraph_format.space_after = Pt(3)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(h_clean)
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(12.5)
                r.font.color.rgb = MAU_MID_BLUE
            elif hashes == 4 or h_clean.startswith("Cặp ") or h_clean.startswith("Bộ "):
                # Tiêu đề cấp 4
                p_h.paragraph_format.space_before = Pt(8)
                p_h.paragraph_format.space_after = Pt(2)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(h_clean)
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = MAU_DARK
            else:
                # Tiêu đề cấp 5
                p_h.paragraph_format.space_before = Pt(6)
                p_h.paragraph_format.space_after = Pt(2)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(h_clean)
                r.bold = True
                r.italic = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = MAU_DARK
                
            i += 1
            continue

        # 9. Danh sách gạch đầu dòng (- hoặc *)
        list_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.+)$', line)
        if list_match:
            indent_spaces = len(list_match.group(1))
            bullet_sym = list_match.group(2)
            content = list_match.group(3)
            
            p_li = doc.add_paragraph()
            p_li.paragraph_format.line_spacing = 1.3
            p_li.paragraph_format.space_before = Pt(1)
            p_li.paragraph_format.space_after = Pt(3)
            
            level = 1 if indent_spaces >= 2 else 0
            p_li.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
            p_li.paragraph_format.first_line_indent = Inches(-0.15)
            
            prefix = "• " if bullet_sym in ["-", "*"] else f"{bullet_sym} "
            r_bullet = p_li.add_run(prefix)
            r_bullet.bold = True
            r_bullet.font.name = "Times New Roman"
            r_bullet.font.size = Pt(11.5)
            r_bullet.font.color.rgb = MAU_MID_BLUE if bullet_sym in ["-", "*"] else MAU_NAVY
            
            format_runs_in_paragraph(p_li, content, default_size=11.5, default_color=MAU_TEXT)
            i += 1
            continue

        # 10. Đoạn văn bản thông thường (Paragraph)
        if stripped:
            p_norm = doc.add_paragraph()
            p_norm.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_norm.paragraph_format.line_spacing = 1.3
            p_norm.paragraph_format.space_before = Pt(2)
            p_norm.paragraph_format.space_after = Pt(5)
            format_runs_in_paragraph(p_norm, stripped, default_size=12, default_color=MAU_TEXT)
            
        i += 1

    if in_table: flush_table()
    if in_callout: flush_callout()
    if in_code_block: flush_code()
    if in_math_block: flush_math()

    print(f"\n>>> Tổng kết chuyển đổi:")
    print(f"  - Số ảnh đã chèn: {images_added}/10")
    print(f"  - Số bảng đã tạo: {tables_added}")
    print(f">>> Đang lưu tài liệu Word tại: {out_path}...")
    doc.save(out_path)
    print(f"[THÀNH CÔNG] Đã tạo file Word hoàn chỉnh tại: {out_path.resolve()}")


def main():
    if not FILE_MD.exists():
        print(f"[LỖI] Không tìm thấy file {FILE_MD}")
        sys.exit(1)
        
    render_markdown_to_docx(FILE_MD, OUT_DOCX_REPORT)
    
    import shutil
    shutil.copy2(OUT_DOCX_REPORT, OUT_DOCX_ROOT)
    print(f"[ĐỒNG BỘ] Đã sao chép sang: {OUT_DOCX_ROOT.resolve()}")


if __name__ == "__main__":
    main()
